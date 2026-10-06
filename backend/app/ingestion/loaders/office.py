"""Office document loader (.docx / .pptx / .xlsx, legacy .doc / .ppt / .xls via LibreOffice).

Content comes from docling: one section per slide ("## Slide N") / sheet ("## Sheet: name"),
which gives citations a location. Author/date/title come from the file's core properties,
falling back to labels in the body ("Author:", "Prepared by:", ...). Files with an Office
extension that are not real OOXML (e.g. text saved as .pptx) are parsed as text and flagged.
"""

import hashlib
import html
import logging
import re
import shutil
import subprocess
import tempfile
from datetime import date as Date
from datetime import datetime
from functools import lru_cache
from pathlib import Path

from app.ingestion.docling_md import normalize_markdown
from app.ingestion.loaders.meeting import _LABEL, make_doc_id, parse_date, parse_people
from app.models.source import SourceDoc

log = logging.getLogger(__name__)

OFFICE_TYPES = {
    ".docx": "docx",
    ".pptx": "pptx",
    ".xlsx": "xlsx",
    ".doc": "docx",
    ".ppt": "pptx",
    ".xls": "xlsx",
}
_LEGACY = {".doc", ".ppt", ".xls"}
_GENERIC_AUTHORS = {
    "",
    "python-docx",
    "python-pptx",
    "openpyxl",
    "unknown",
    "author",
    "user",
    "microsoft office user",
    "admin",
    "administrator",
}
_AUTHOR_RE = re.compile(
    _LABEL.format(
        labels="author|authors|owner|document owner|prepared by|presented by|"
        "presenter|submitted by|written by|created by"
    ),
    re.I | re.M,
)
_REVIEWER_RE = re.compile(
    _LABEL.format(labels="reviewed by|reviewers?|approved by|approvers?"), re.I | re.M
)
_TITLE_RE = re.compile(_LABEL.format(labels="title|document title|subject"), re.I | re.M)
_GENERIC_TITLES = {
    "",
    "word document",
    "powerpoint presentation",
    "excel workbook",
    "document",
    "presentation",
    "workbook",
    "untitled",
    "title",
}
# Creation timestamps baked into the default templates of python-docx / python-pptx: a file
# generated from them carries these dates, which say nothing about the document.
_TEMPLATE_TIMESTAMPS = {datetime(2013, 12, 23, 23, 15), datetime(2013, 1, 27, 9, 14, 16)}
_DATE_RE = re.compile(_LABEL.format(labels="date|report date|revision date|as of"), re.I | re.M)


def parse_byline(value: str) -> dict[str, str | None]:
    """'James Ortiz, Lead Yield Engineer, PE Division' -> {'James Ortiz': 'Lead Yield ...'}.
    Several people may be separated by '|' or ';'. Parenthesised roles are also handled."""
    people: dict[str, str | None] = {}
    for part in re.split(r"\s*[|;]\s*", value):
        if not part.strip():
            continue
        if "(" in part or " and " in part or "&" in part:
            people |= parse_people(part)
            continue
        name, _, role = part.partition(",")
        people |= {n: (role.strip(" .") or r) for n, r in parse_people(name).items()}
    return people


_HEADING = re.compile(r"^#{1,6}\s+(.+?)\s*$", re.M)


def iter_document_files(documents_dir: Path) -> list[Path]:
    if not documents_dir.is_dir():
        return []
    return sorted(
        p
        for p in documents_dir.rglob("*")
        if p.is_file() and p.suffix.lower() in OFFICE_TYPES and not p.name.startswith(("~$", "."))
    )


def is_ooxml(path: Path) -> bool:
    with path.open("rb") as f:
        return f.read(4) == b"PK\x03\x04"


def _core_properties(path: Path, kind: str) -> dict:
    try:
        if kind == "docx":
            from docx import Document

            cp = Document(str(path)).core_properties
            return {
                "author": cp.author or cp.last_modified_by,
                "title": cp.title,
                "created": cp.created or cp.modified,
            }
        if kind == "pptx":
            from pptx import Presentation

            cp = Presentation(str(path)).core_properties
            return {
                "author": cp.author or cp.last_modified_by,
                "title": cp.title,
                "created": cp.created or cp.modified,
            }
        from openpyxl import load_workbook

        props = load_workbook(str(path), read_only=True).properties
        return {
            "author": props.creator or props.lastModifiedBy,
            "title": props.title,
            "created": props.created or props.modified,
        }
    except Exception:  # corrupt / unusual files: fall back to body parsing
        log.warning("could not read core properties of %s", path, exc_info=True)
        return {}


def _demote(md: str, levels: int = 2) -> str:
    """Push headings inside a slide/sheet below the '## Slide N' section heading."""
    return re.sub(
        r"^(#{1,6})\s", lambda m: "#" * min(6, len(m.group(1)) + levels) + " ", md, flags=re.M
    )


@lru_cache(maxsize=1)
def _converter():
    from docling.datamodel.base_models import InputFormat
    from docling.document_converter import DocumentConverter

    return DocumentConverter(allowed_formats=[InputFormat.DOCX, InputFormat.PPTX, InputFormat.XLSX])


def _slide_notes(path: Path) -> dict[int, str]:
    from pptx import Presentation

    notes = {}
    for i, slide in enumerate(Presentation(str(path)).slides, start=1):
        if slide.has_notes_slide:
            text = slide.notes_slide.notes_text_frame.text.strip()
            if text:
                notes[i] = text
    return notes


def _sheet_as_of(path: Path) -> str | None:
    """Latest date in a sheet's date columns (not due/target dates): the data's "as of" date."""
    from openpyxl import load_workbook

    latest: Date | None = None
    for ws in load_workbook(str(path), read_only=True, data_only=True).worksheets:
        rows = ws.iter_rows(values_only=True)
        header = [str(h or "").lower() for h in next(rows, ())]
        cols = [
            i
            for i, h in enumerate(header)
            if "date" in h and not any(w in h for w in ("due", "target", "planned"))
        ]
        for row in rows:
            for i in cols:
                iso = parse_date(row[i]) if i < len(row) and row[i] not in (None, "") else None
                if iso and (latest is None or Date.fromisoformat(iso) > latest):
                    latest = Date.fromisoformat(iso)
    return latest.isoformat() if latest else None


def _sheet_names(path: Path) -> list[str]:
    from openpyxl import load_workbook

    return load_workbook(str(path), read_only=True).sheetnames


def docling_markdown(path: Path, kind: str) -> tuple[str, int]:
    """Return (markdown, number of slides/sheets/pages)."""
    doc = _converter().convert(str(path)).document
    pages = sorted(doc.pages)
    if kind == "pptx" and pages:
        notes = _slide_notes(path)
        parts = []
        for n in pages:
            body = _demote(doc.export_to_markdown(page_no=n).strip())
            if n in notes:
                body += f"\n\nSpeaker notes: {notes[n]}"
            parts.append(f"## Slide {n}\n\n{body}")
        return "\n\n".join(parts), len(pages)
    if kind == "xlsx" and pages:
        names = _sheet_names(path)
        parts = []
        for i, n in enumerate(pages):
            name = names[i] if i < len(names) else f"Sheet {n}"
            parts.append(
                f"## Sheet: {name}\n\n{_demote(doc.export_to_markdown(page_no=n).strip())}"
            )
        return "\n\n".join(parts), len(pages)
    return doc.export_to_markdown(), max(1, len(pages))


def _convert_legacy(path: Path, kind: str, out_dir: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    subprocess.run(
        [soffice, "--headless", "--convert-to", kind, "--outdir", str(out_dir), str(path)],
        check=True,
        capture_output=True,
        timeout=180,
    )
    converted = out_dir / f"{path.stem}.{kind}"
    return converted if converted.exists() else None


def _title_from(md: str, fallback: str) -> str:
    if m := _HEADING.search(md):
        title = re.sub(r"[*_`]", "", m.group(1)).strip()
        if title and not re.fullmatch(r"(Slide \d+|Sheet: .*)", title):
            return title
    return fallback.replace("_", " ").replace("-", " ").strip().title()


def load_office(path: Path, data_dir: Path) -> SourceDoc:
    rel = path.relative_to(data_dir).as_posix()
    raw = path.read_bytes()
    kind = OFFICE_TYPES[path.suffix.lower()]
    warnings: list[str] = []
    pages = 1

    with tempfile.TemporaryDirectory() as tmp:
        target = path
        if path.suffix.lower() in _LEGACY:
            converted = _convert_legacy(path, kind, Path(tmp))
            if not converted:
                raise RuntimeError("legacy Office format needs LibreOffice (soffice) installed")
            target = converted

        if is_ooxml(target):
            props = _core_properties(target, kind)
            if kind == "xlsx":
                props["as_of"] = _sheet_as_of(target)
            content, pages = docling_markdown(target, kind)
            content = html.unescape(content)
            fmt = "ooxml"
        else:
            props = {}
            text = raw.decode("utf-8", errors="replace")
            content = normalize_markdown(text, f"{path.stem}.md")
            fmt = "text-fallback"
            warnings.append(f"{rel}: not a real {kind} file (plain text); parsed as text")

    head = content[:4000]
    author = (props.get("author") or "").strip()
    authors: dict[str, str | None] = {}
    people_source = "none"
    if author.lower() not in _GENERIC_AUTHORS:
        authors, people_source = parse_byline(author), "core_properties"
    elif m := _AUTHOR_RE.search(head):
        authors, people_source = parse_byline(m.group("value")), "body"
    reviewers = parse_byline(m.group("value")) if (m := _REVIEWER_RE.search(head)) else {}

    created = props.get("created")
    trusted_created = (
        isinstance(created, datetime)
        and created.replace(tzinfo=None, microsecond=0) not in _TEMPLATE_TIMESTAMPS
        and author.lower() not in _GENERIC_AUTHORS  # written by a library: it's generation time
    )
    date = parse_date(created) if trusted_created else None
    date_basis = "core_properties" if date else None
    if m := _DATE_RE.search(content[:4000]):
        if body_date := parse_date(m.group("value")):  # an explicit date in the body wins
            date, date_basis = body_date, "body"
    if not date and (as_of := props.get("as_of")):
        date, date_basis = as_of, "latest_row_date"

    title = (props.get("title") or "").strip()
    if title.lower() in _GENERIC_TITLES:
        title = ""
    if not title and (m := _TITLE_RE.search(head)):
        title = re.sub(r"[*_`]", "", m.group("value")).strip()
    title = title or _title_from(content, path.stem)
    extra = {"format": fmt, "pages": pages, "warnings": warnings, "date_basis": date_basis}
    if roles := {n: r for n, r in authors.items() if r}:
        extra["author_roles"] = roles
    if reviewers:
        extra["reviewers"] = list(reviewers)
        extra["author_roles"] = {
            **{n: r for n, r in reviewers.items() if r},
            **extra.get("author_roles", {}),
        }

    return SourceDoc(
        doc_id=make_doc_id(rel),
        source_file=rel,
        source_type=kind,
        title=title,
        date=date,
        authors=list(authors),
        attendees_source=people_source,
        content=content,
        content_hash=hashlib.sha256(raw).hexdigest(),
        extra=extra,
    )
