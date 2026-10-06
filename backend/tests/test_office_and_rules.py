from pathlib import Path

from docx import Document as Docx
from openpyxl import Workbook
from pptx import Presentation

from app.ingest import count_tokens, split_table
from app.loaders import iter_document_files, load_office
from app.rules import apply_ingest_rules
from app.schemas import ActionItem, EnrichmentResult, SourceDoc


def _mock_office_files(docs: Path) -> None:
    """Small but real OOXML files, plus a non-Office file that must be ignored."""
    (docs / "docx").mkdir(parents=True)
    (docs / "pptx").mkdir()
    (docs / "xlsx").mkdir()
    d = Docx()
    d.core_properties.author = "Anna Becker"
    d.add_heading("8D Corrective Action", 1)
    d.add_paragraph("Containment for field returns on lot 4412.")
    d.save(docs / "docx" / "corrective_action_8d.docx")

    p = Presentation()
    p.core_properties.author = "Lisa Park"
    for i in (1, 2):
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = f"Ramp status {i}"
        s.placeholders[1].text = f"Gate {i} on track"
        s.notes_slide.notes_text_frame.text = f"note {i}"
    p.save(docs / "pptx" / "ramp_status.pptx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Yield"
    ws.append(["Lot", "Yield"])
    ws.append(["LOT-1", 0.41])
    wb.properties.creator = "James Ortiz"
    wb.save(docs / "xlsx" / "yield_tracker.xlsx")

    (docs / "pptx" / "text_deck.pptx").write_text("# Audit Readout\n\nPresenter: Mike Chen\n")
    (docs / "pptx" / "text_deck_slides.md").write_text("ignored: not an office extension")


def test_office_loader_real_and_text_fallback(isolated_env):
    docs = isolated_env / "documents"
    _mock_office_files(docs)
    files = {p.name: p for p in iter_document_files(docs)}
    assert set(files) == {
        "corrective_action_8d.docx",
        "ramp_status.pptx",
        "yield_tracker.xlsx",
        "text_deck.pptx",
    }

    docx = load_office(files["corrective_action_8d.docx"], isolated_env)
    assert (docx.source_type, docx.authors, docx.title) == (
        "docx",
        ["Anna Becker"],
        "8D Corrective Action",
    )

    pptx = load_office(files["ramp_status.pptx"], isolated_env)
    assert pptx.authors == ["Lisa Park"] and pptx.extra["pages"] == 2
    assert "## Slide 2" in pptx.content and "Speaker notes: note 2" in pptx.content

    xlsx = load_office(files["yield_tracker.xlsx"], isolated_env)
    assert xlsx.authors == ["James Ortiz"] and "## Sheet: Yield" in xlsx.content

    text = load_office(files["text_deck.pptx"], isolated_env)
    assert text.extra["format"] == "text-fallback" and text.extra["warnings"]
    assert text.source_file == "documents/pptx/text_deck.pptx"
    assert text.authors == ["Mike Chen"] and text.attendees_source == "body"


def test_business_rules():
    src = SourceDoc(
        doc_id="d",
        source_file="f",
        source_type="meeting",
        title="t",
        content="NovaDrive line-down risk. Lisa Park owns containment for Volta-7.",
        content_hash="h",
    )
    enr = EnrichmentResult(
        topic_domain="customer",
        priority="medium",
        products=["Volta-7", "Falcon-7"],
        summary="",
        key_topics=[],
        decisions=[],
        action_items=[
            ActionItem(owner="Lisa Park", task="containment", due_date=None),
            ActionItem(owner="Ghost Person", task="x", due_date=None),
        ],
    )
    out, applied = apply_ingest_rules(src, enr)
    assert out.priority == "critical"
    assert [a.owner for a in out.action_items] == ["Lisa Park", "unassigned"]
    assert out.products == ["Volta-7"]
    assert [a.split(":")[0] for a in applied] == ["R1", "R2", "R3"]


def test_table_split_repeats_header():
    rows = "\n".join(f"| LOT-{i} | {i}% |" for i in range(200))
    pieces = split_table(f"| Lot | Yield |\n|---|---|\n{rows}", 100)
    assert len(pieces) > 1
    assert all(p.startswith("| Lot | Yield |\n|---|---|") for p in pieces)
    assert all(count_tokens(p) <= 100 for p in pieces)


def test_sheet_as_of_ignores_columns_that_only_contain_date_letters(tmp_path: Path):
    from app.loaders import _sheet_as_of

    wb = Workbook()
    ws = wb.active
    ws.append(["Date", "Status Update", "Validated By", "Due_Date"])
    ws.append(["2026-03-02", "moved to week 12", "2026-12-31", "2026-11-30"])
    path = tmp_path / "t.xlsx"
    wb.save(path)
    assert _sheet_as_of(path) == "2026-03-02"


def test_sheet_as_of_accepts_date_header_variants(tmp_path: Path):
    from app.loaders import _sheet_as_of

    for header in ["Date:", "Dates", "Date(UTC)", "Log Date.", "Date_Raised", "closed_date"]:
        wb = Workbook()
        wb.active.append([header, "Status Update"])
        wb.active.append(["2026-04-02", "2026-12-31"])
        path = tmp_path / "h.xlsx"
        wb.save(path)
        assert _sheet_as_of(path) == "2026-04-02", header
